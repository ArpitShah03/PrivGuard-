import os
import tempfile
import pytest
import docx
import pptx
import fitz

from models.schemas import DocumentBlock, PIIFinding
from redaction.document_sanitizer import (
    sanitize_docx,
    sanitize_pptx,
    sanitize_pdf,
    sanitize_txt,
    DocumentSanitizer
)
from extractors.router import route_and_extract
from detection.pii_engine import detect_pii_in_document_blocks
from detection.entity_resolver import EntityRegistry, EntityResolver


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmp:
        yield tmp


# 1. DOCX PII replacement
def test_docx_pii_replacement(temp_dir):
    input_docx = os.path.join(temp_dir, "test.docx")
    output_docx = os.path.join(temp_dir, "sanitized_test.docx")

    doc = docx.Document()
    doc.add_paragraph("Report prepared by Tobias Lindqvist.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Director"
    table.cell(0, 1).text = "Tobias Lindqvist"
    table.cell(1, 0).text = "Email"
    table.cell(1, 1).text = "tlindqvist@cadence.com"
    doc.save(input_docx)

    findings = [
        PIIFinding("1", "d", "test.docx", "paragraph_1", 19, 35, "Tobias Lindqvist", "PERSON", 0.95, ["NER"], "HIGH", replacement_token="[PERSON_001]"),
        PIIFinding("2", "d", "test.docx", "table_1_row_1_cell_2", 0, 16, "Tobias Lindqvist", "PERSON", 0.95, ["NER"], "HIGH", replacement_token="[PERSON_001]"),
        PIIFinding("3", "d", "test.docx", "table_1_row_2_cell_2", 0, 22, "tlindqvist@cadence.com", "EMAIL", 0.95, ["REGEX"], "HIGH", replacement_token="[EMAIL_001]")
    ]

    sanitize_docx(input_docx, output_docx, findings)

    # Re-open and verify
    reopened = docx.Document(output_docx)
    assert "[PERSON_001]" in reopened.paragraphs[0].text
    assert "Tobias Lindqvist" not in reopened.paragraphs[0].text
    assert "[PERSON_001]" in reopened.tables[0].cell(0, 1).text
    assert "[EMAIL_001]" in reopened.tables[0].cell(1, 1).text


# 2. PPTX PII replacement
def test_pptx_pii_replacement(temp_dir):
    input_pptx = os.path.join(temp_dir, "test.pptx")
    output_pptx = os.path.join(temp_dir, "sanitized_test.pptx")

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(100, 100, 300, 100)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Lead Architect: Sarah Connor (sarah.connor@cadence.com)"
    prs.save(input_pptx)

    findings = [
        PIIFinding("1", "d", "test.pptx", "s1", 16, 28, "Sarah Connor", "PERSON", 0.95, ["NER"], "HIGH", replacement_token="[PERSON_002]"),
        PIIFinding("2", "d", "test.pptx", "s1", 30, 54, "sarah.connor@cadence.com", "EMAIL", 0.95, ["REGEX"], "HIGH", replacement_token="[EMAIL_002]")
    ]

    sanitize_pptx(input_pptx, output_pptx, findings)

    reopened = pptx.Presentation(output_pptx)
    slide_text = reopened.slides[0].shapes[0].text_frame.text
    assert "[PERSON_002]" in slide_text
    assert "[EMAIL_002]" in slide_text
    assert "Sarah Connor" not in slide_text
    assert "sarah.connor@cadence.com" not in slide_text


# 3. PDF native-text redaction
def test_pdf_native_text_redaction(temp_dir):
    input_pdf = os.path.join(temp_dir, "test.pdf")
    output_pdf = os.path.join(temp_dir, "sanitized_test.pdf")

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 100), "Confidential applicant: Desmond Achterberg")
    doc.save(input_pdf)
    doc.close()

    findings = [
        PIIFinding("1", "d", "test.pdf", "page_1", 24, 42, "Desmond Achterberg", "PERSON", 0.95, ["NER"], "HIGH", replacement_token="[PERSON_003]")
    ]

    sanitize_pdf(input_pdf, output_pdf, findings)

    # Re-open and verify text is sanitized in PDF stream
    reopened = fitz.open(output_pdf)
    text = reopened[0].get_text()
    reopened.close()
    assert "[PERSON_003]" in text
    assert "Desmond Achterberg" not in text


# 4. TXT redaction
def test_txt_redaction(temp_dir):
    input_txt = os.path.join(temp_dir, "test.txt")
    output_txt = os.path.join(temp_dir, "sanitized_test.txt")

    with open(input_txt, "w", encoding="utf-8") as f:
        f.write("Citizen SSN: 219-45-7821 on file.")

    findings = [
        PIIFinding("1", "d", "test.txt", "p1", 13, 24, "219-45-7821", "SSN", 0.98, ["REGEX"], "CRITICAL", replacement_token="[SSN_001]")
    ]

    sanitize_txt(input_txt, output_txt, findings)

    with open(output_txt, "r", encoding="utf-8") as f:
        sanitized = f.read()

    assert "[SSN_001]" in sanitized
    assert "219-45-7821" not in sanitized


# 5. Sanitized output can be reopened
def test_sanitized_output_can_be_reopened(temp_dir):
    # DOCX
    docx_in = os.path.join(temp_dir, "a.docx")
    docx_out = os.path.join(temp_dir, "a_san.docx")
    d = docx.Document()
    d.add_paragraph("Clean sample")
    d.save(docx_in)
    sanitize_docx(docx_in, docx_out, [])
    assert docx.Document(docx_out) is not None

    # PPTX
    pptx_in = os.path.join(temp_dir, "b.pptx")
    pptx_out = os.path.join(temp_dir, "b_san.pptx")
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[6])
    p.save(pptx_in)
    sanitize_pptx(pptx_in, pptx_out, [])
    assert pptx.Presentation(pptx_out) is not None

    # PDF
    pdf_in = os.path.join(temp_dir, "c.pdf")
    pdf_out = os.path.join(temp_dir, "c_san.pdf")
    pdf_doc = fitz.open()
    pdf_doc.new_page()
    pdf_doc.save(pdf_in)
    pdf_doc.close()
    sanitize_pdf(pdf_in, pdf_out, [])
    reopened_pdf = fitz.open(pdf_out)
    assert len(reopened_pdf) == 1
    reopened_pdf.close()


# 6. Residual PII causes BLOCKED
def test_residual_pii_causes_blocked(temp_dir):
    sanitizer = DocumentSanitizer(output_dir=temp_dir)
    txt_file = os.path.join(temp_dir, "leaked.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        # Deliberately contain an unredacted valid SSN
        f.write("Residual SSN: 219-45-7821 remaining in output.")

    orig_block = DocumentBlock("d", "leaked.txt", "txt", "p1", "text", "Orig text", 0)
    val, _ = sanitizer.verify_and_rescan(txt_file, [orig_block], [], "d")

    assert val.is_safe is False
    assert val.requires_human_review is False
    assert val.security_decision == "BLOCKED"
    assert val.high_risk_residual >= 1


# 7. Extraction uncertainty causes HUMAN_REVIEW
def test_extraction_uncertainty_causes_human_review(temp_dir):
    sanitizer = DocumentSanitizer(output_dir=temp_dir)
    txt_file = os.path.join(temp_dir, "clean.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write("All clear content.")

    # Original document had an uncertain scanned block
    uncertain_block = DocumentBlock(
        "d", "scan.pdf", "pdf", "page_1", "OCR_UNAVAILABLE",
        "[OCR_UNAVAILABLE: Tesseract missing]", 0,
        metadata={"is_uncertain": True}
    )

    val, _ = sanitizer.verify_and_rescan(txt_file, [uncertain_block], [], "d")

    assert val.is_safe is False
    assert val.requires_human_review is True
    assert val.security_decision == "HUMAN_REVIEW"
    assert val.uncertain_blocks_count >= 1


# 8. Replacement tokens do not trigger false PII
def test_replacement_tokens_do_not_trigger_false_pii(temp_dir):
    sanitizer = DocumentSanitizer(output_dir=temp_dir)
    txt_file = os.path.join(temp_dir, "tokens.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write("Contact [PERSON_001] via [EMAIL_001] with SSN [SSN_001].")

    clean_block = DocumentBlock("d", "tokens.txt", "txt", "p1", "text", "Clean text", 0)
    val, re_blocks = sanitizer.verify_and_rescan(txt_file, [clean_block], [], "d")

    # Tokens must be ignored; zero residual findings
    assert len(val.remaining_findings) == 0
    assert val.is_safe is True
    assert val.security_decision == "SANITIZED"


# 9. No double-redaction
def test_no_double_redaction(temp_dir):
    txt_in = os.path.join(temp_dir, "double.txt")
    txt_out = os.path.join(temp_dir, "double_san.txt")
    with open(txt_in, "w", encoding="utf-8") as f:
        f.write("Manager: Alice Brown.")

    finding = PIIFinding("1", "d", "double.txt", "p1", 9, 20, "Alice Brown", "PERSON", 0.95, ["NER"], "HIGH", replacement_token="[PERSON_005]")

    # Redact once
    sanitize_txt(txt_in, txt_out, [finding])
    with open(txt_out, "r", encoding="utf-8") as f:
        first_pass = f.read()

    # Redact twice
    sanitize_txt(txt_out, txt_out, [finding])
    with open(txt_out, "r", encoding="utf-8") as f:
        second_pass = f.read()

    assert first_pass == second_pass
    assert "[[PERSON_005]" not in second_pass
    assert "[PERSON_005]" in second_pass


# 10. End-to-end DETECT -> REDACT -> RESCAN
def test_end_to_end_detect_redact_rescan(temp_dir):
    sanitizer = DocumentSanitizer(output_dir=temp_dir)
    txt_in = os.path.join(temp_dir, "employee.txt")
    with open(txt_in, "w", encoding="utf-8") as f:
        f.write("Applicant: Tobias Lindqvist, Email: tlindqvist@cadence.com, SSN: 219-45-7821.")

    # 1. Extract
    blocks = route_and_extract(txt_in, "doc_e2e", "employee.txt")
    assert len(blocks) >= 1

    # 2. Detect & Resolve
    findings = detect_pii_in_document_blocks(blocks)
    assert len(findings) >= 2

    # 3. Redact
    sanitized_path = sanitizer.sanitize(txt_in, findings, "employee_sanitized.txt")
    assert os.path.exists(sanitized_path)

    # 4. Verify & Re-Scan
    val, re_blocks = sanitizer.verify_and_rescan(sanitized_path, blocks, findings, "doc_e2e")

    # Clean document with zero uncertainty and zero residual PII
    assert val.is_safe is True
    assert val.requires_human_review is False
    assert val.security_decision == "SANITIZED"
    assert val.high_risk_residual == 0
