import pytest
import io
import fitz
import docx
import pptx
from pptx.util import Inches
from PIL import Image

from models.schemas import DocumentBlock, PIIFinding, DocumentAnalysisResult
from extractors.ocr_engine import OCREngine, OCRResult, OCRStatus
from extractors.pdf_extractor import extract_pdf, is_page_scanned
from extractors.docx_extractor import extract_docx
from extractors.pptx_extractor import extract_pptx
from extractors.txt_extractor import extract_txt
from extractors.metadata_extractor import extract_metadata_as_blocks
from extractors.router import route_and_extract, compute_extraction_metrics


# ====================================================================
# 1. Schema Backward Compatibility Tests
# ====================================================================

def test_schema_document_block_backward_compatibility():
    """Verify legacy 7 positional arguments constructor works without error."""
    block = DocumentBlock(
        "doc1", "test.pdf", "pdf", "page_1", "text", "Sample content", 0
    )
    assert block.document_id == "doc1"
    assert block.bbox is None
    assert block.is_ocr is False
    assert block.ocr_confidence == 1.0
    assert block.metadata == {}
    assert block.parent_id is None


def test_schema_document_block_extended_fields():
    """Verify new Phase 1 fields initialize properly."""
    block = DocumentBlock(
        document_id="doc2",
        file_name="scan.pdf",
        file_type="pdf",
        location="page_1_ocr_block_1",
        block_type="ocr_text",
        text="Scanned text",
        order=1,
        bbox={"x0": 0.1, "y0": 0.2, "x1": 0.8, "y1": 0.9},
        is_ocr=True,
        ocr_confidence=0.92,
        metadata={"source": "ocr"},
        parent_id="page_1"
    )
    assert block.is_ocr is True
    assert block.ocr_confidence == 0.92
    assert block.bbox["x0"] == 0.1
    assert block.parent_id == "page_1"


def test_schema_pii_finding_backward_compatibility():
    """Verify PIIFinding legacy constructor works and has new optional defaults."""
    finding = PIIFinding(
        "f1", "d1", "test.txt", "para_1", 0, 10, "John Smith", "PERSON", 0.95, ["NER"], "HIGH"
    )
    assert finding.finding_id == "f1"
    assert finding.bbox is None
    assert finding.is_ocr is False
    assert finding.replacement_token is None


def test_schema_document_analysis_result_extended():
    """Verify DocumentAnalysisResult supports HUMAN_REVIEW and extraction metadata."""
    res = DocumentAnalysisResult(
        document_id="d1",
        file_name="doc.pdf",
        file_type="pdf",
        status="HUMAN_REVIEW",
        original_blocks=[],
        sanitized_blocks=[],
        findings=[],
        extraction_metadata={"total_blocks": 5},
        ocr_coverage=0.85
    )
    assert res.status == "HUMAN_REVIEW"
    assert res.ocr_coverage == 0.85
    assert res.extraction_metadata["total_blocks"] == 5


# ====================================================================
# 2. OCR Engine Tests (Unit & Mocked)
# ====================================================================

class MockOCREngine(OCREngine):
    """Custom mock engine to test different OCR statuses deterministically."""
    def __init__(self, forced_status: str, forced_text: str = "", forced_conf: float = 0.0):
        super().__init__()
        self.forced_status = forced_status
        self.forced_text = forced_text
        self.forced_conf = forced_conf

    @property
    def is_available(self) -> bool:
        return self.forced_status != OCRStatus.UNAVAILABLE

    def extract_text(self, image_input, lang="eng", preprocess=True) -> OCRResult:
        if self.forced_status == OCRStatus.UNAVAILABLE:
            return OCRResult(
                status=OCRStatus.UNAVAILABLE,
                text="[OCR_UNAVAILABLE]",
                confidence=0.0,
                error_message="Tesseract binary not found",
                is_uncertain=True
            )
        elif self.forced_status == OCRStatus.FAILED:
            return OCRResult(
                status=OCRStatus.FAILED,
                text="[OCR_FAILED]",
                confidence=0.0,
                error_message="Engine crashed",
                is_uncertain=True
            )
        elif self.forced_status == OCRStatus.LOW_CONFIDENCE:
            return OCRResult(
                status=OCRStatus.LOW_CONFIDENCE,
                text=self.forced_text or "corrupted ?#! text",
                confidence=self.forced_conf or 0.35,
                blocks=[{"text": self.forced_text or "corrupted", "confidence": 0.35, "bbox": {"x0": 0, "y0": 0, "x1": 1, "y1": 1}}],
                is_uncertain=True
            )
        return OCRResult(
            status=OCRStatus.SUCCESS,
            text=self.forced_text or "Successfully Extracted OCR Text",
            confidence=self.forced_conf or 0.95,
            blocks=[{"text": self.forced_text or "Successfully Extracted OCR Text", "confidence": 0.95, "bbox": {"x0": 0, "y0": 0, "x1": 1, "y1": 1}}],
            is_uncertain=False
        )


def test_ocr_engine_preprocessing():
    """Verify Pillow image preprocessing (grayscale and contrast)."""
    engine = OCREngine()
    # Create simple color image
    img = Image.new("RGB", (100, 100), color=(255, 0, 0))
    processed = engine.preprocess_image(img, grayscale=True, enhance_contrast=True)
    assert processed.mode == "L"
    assert processed.size == (100, 100)


def test_ocr_engine_unavailable_handling():
    """Verify OCR engine explicitly marks uncertainty when binary is unavailable."""
    engine = MockOCREngine(forced_status=OCRStatus.UNAVAILABLE)
    res = engine.extract_text(b"fake_image_bytes")
    assert res.status == OCRStatus.UNAVAILABLE
    assert res.is_uncertain is True
    assert res.confidence == 0.0


# ====================================================================
# 3. PDF Multimodal Extraction Tests
# ====================================================================

def test_pdf_native_text_and_bbox(tmp_path):
    """Test native PDF extraction with spatial bounding boxes and metadata."""
    pdf_path = tmp_path / "native_sample.pdf"
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    page.insert_text((72, 100), "Confidential Strategic Plan 2026", fontsize=14)
    page.insert_text((72, 150), "Author: Johnathan Vance, Risk Director", fontsize=12)
    doc.set_metadata({
        "title": "Risk Policy",
        "author": "Alice Cooper",
        "subject": "Enterprise Risk"
    })
    doc.save(str(pdf_path))
    doc.close()

    blocks = extract_pdf(str(pdf_path), "doc_pdf_1", "native_sample.pdf")
    assert len(blocks) >= 2

    # Check native text blocks
    native_blocks = [b for b in blocks if b.block_type == "text"]
    assert len(native_blocks) >= 1
    assert any("Confidential Strategic Plan" in b.text for b in native_blocks)
    # Check bounding box
    first_block = native_blocks[0]
    assert first_block.bbox is not None
    assert "x0" in first_block.bbox and "y1" in first_block.bbox
    assert 0.0 <= first_block.bbox["x0"] <= 1.0

    # Check metadata blocks
    meta_blocks = [b for b in blocks if b.block_type == "metadata"]
    assert len(meta_blocks) >= 2
    assert any("Alice Cooper" in b.text for b in meta_blocks)


def test_pdf_scanned_page_detection_and_ocr(tmp_path):
    """Test scanned page detection and routing through OCR engine."""
    pdf_path = tmp_path / "scanned_sample.pdf"
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    
    # Insert visual image with NO native text
    img = Image.new("RGB", (200, 200), color=(200, 200, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(str(pdf_path))
    doc.close()

    # Provide mock engine that succeeds
    mock_engine = MockOCREngine(forced_status=OCRStatus.SUCCESS, forced_text="Scanned Vendor Invoice Data")
    blocks = extract_pdf(str(pdf_path), "doc_scan_1", "scanned_sample.pdf", ocr_engine=mock_engine)

    ocr_blocks = [b for b in blocks if b.is_ocr]
    assert len(ocr_blocks) >= 1
    assert any("Scanned Vendor Invoice Data" in b.text for b in ocr_blocks)
    assert ocr_blocks[0].ocr_confidence == 0.95


def test_pdf_scanned_page_ocr_failure_handling(tmp_path):
    """Test fail-closed behavior when OCR is unavailable on scanned PDF."""
    pdf_path = tmp_path / "scanned_fail.pdf"
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    img = Image.new("RGB", (100, 100), color=(128, 128, 128))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(str(pdf_path))
    doc.close()

    mock_engine = MockOCREngine(forced_status=OCRStatus.UNAVAILABLE)
    blocks = extract_pdf(str(pdf_path), "doc_fail_1", "scanned_fail.pdf", ocr_engine=mock_engine)

    # Must preserve explicit uncertainty block
    uncertain_blocks = [b for b in blocks if b.metadata.get("is_uncertain", False)]
    assert len(uncertain_blocks) == 1
    assert "OCR_UNAVAILABLE" in uncertain_blocks[0].block_type


# ====================================================================
# 4. DOCX Multimodal Extraction Tests
# ====================================================================

def test_docx_complete_multimodal(tmp_path):
    """Test DOCX extraction: paragraphs, tables, headers, footers, metadata."""
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()

    # Paragraphs
    doc.add_paragraph("First introductory paragraph.")
    doc.add_paragraph("Second policy paragraph.")

    # Table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Vendor Name"
    table.cell(0, 1).text = "Risk Level"
    table.cell(1, 0).text = "Acme Corp"
    table.cell(1, 1).text = "High"

    # Header and Footer
    section = doc.sections[0]
    section.header.paragraphs[0].text = "CONFIDENTIAL INTERNAL USE ONLY"
    section.footer.paragraphs[0].text = "Page 1 - Cadence Financial Group"

    # Core properties
    doc.core_properties.author = "Auditor James Lee"
    doc.core_properties.comments = "Reviewed by compliance"

    doc.save(str(docx_path))

    blocks = extract_docx(str(docx_path), "doc_docx_1", "sample.docx")

    # 1. Check header & footer
    headers = [b for b in blocks if b.block_type == "header"]
    footers = [b for b in blocks if b.block_type == "footer"]
    assert len(headers) == 1
    assert "CONFIDENTIAL INTERNAL USE ONLY" in headers[0].text
    assert len(footers) == 1
    assert "Cadence Financial Group" in footers[0].text

    # 2. Check paragraphs
    paras = [b for b in blocks if b.block_type == "text"]
    assert len(paras) == 2
    assert paras[0].text == "First introductory paragraph."

    # 3. Check table cells with structured coordinates
    cells = [b for b in blocks if b.block_type == "table_cell"]
    assert len(cells) == 4
    acme_cell = next(c for c in cells if "Acme Corp" in c.text)
    assert acme_cell.metadata["row_index"] == 2
    assert acme_cell.metadata["col_index"] == 1

    # 4. Check metadata
    metas = [b for b in blocks if b.block_type == "metadata"]
    assert any("Auditor James Lee" in m.text for m in metas)


# ====================================================================
# 5. PPTX Multimodal Extraction Tests
# ====================================================================

def test_pptx_complete_multimodal(tmp_path):
    """Test PPTX extraction: text shapes, tables, notes, metadata."""
    pptx_path = tmp_path / "presentation.pptx"
    prs = pptx.Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide

    # Normal text shape
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    txBox.text_frame.text = "Executive Leadership Overview"

    # Table
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2.5), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text_frame.text = "Leader Name"
    table.cell(0, 1).text_frame.text = "Title"
    table.cell(1, 0).text_frame.text = "Sarah Connor"
    table.cell(1, 1).text_frame.text = "CISO"

    # Speaker notes
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "Confidential speaker note: Do not distribute."

    # Core properties
    prs.core_properties.author = "VP Operations"
    prs.save(str(pptx_path))

    blocks = extract_pptx(str(pptx_path), "doc_pptx_1", "presentation.pptx")

    # Check text shapes
    shape_blocks = [b for b in blocks if b.block_type == "text"]
    assert any("Executive Leadership Overview" in b.text for b in shape_blocks)

    # Check table cells
    cell_blocks = [b for b in blocks if b.block_type == "table_cell"]
    assert any("Sarah Connor" in b.text for b in cell_blocks)

    # Check speaker notes
    note_blocks = [b for b in blocks if b.block_type == "notes"]
    assert len(note_blocks) == 1
    assert "Confidential speaker note" in note_blocks[0].text
    assert note_blocks[0].location == "slide_1_notes"

    # Check metadata
    meta_blocks = [b for b in blocks if b.block_type == "metadata"]
    assert any("VP Operations" in m.text for m in meta_blocks)


# ====================================================================
# 6. TXT Resilient Fallback Tests
# ====================================================================

def test_txt_encoding_fallback(tmp_path):
    """Test TXT extraction with non-UTF8 (Latin-1) encoding."""
    txt_path = tmp_path / "latin1.txt"
    # Write Latin-1 accented bytes
    latin1_content = "Résumé of John Müller\n\nEmployee ID: EMP-4421".encode("latin-1")
    txt_path.write_bytes(latin1_content)

    blocks = extract_txt(str(txt_path), "doc_txt_1", "latin1.txt")
    assert len(blocks) == 2
    assert "John Müller" in blocks[0].text
    assert blocks[0].location == "paragraph_1"
    assert blocks[1].location == "paragraph_2"


# ====================================================================
# 7. Router & Metrics Tests
# ====================================================================

def test_router_dispatch_and_metrics(tmp_path):
    """Test central router and extraction metrics computation."""
    txt_path = tmp_path / "route_test.txt"
    txt_path.write_text("Hello World\n\nParagraph 2", encoding="utf-8")

    blocks = route_and_extract(str(txt_path), "doc_route_1", "route_test.txt")
    assert len(blocks) == 2

    metrics = compute_extraction_metrics(blocks)
    assert metrics["total_blocks"] == 2
    assert metrics["native_blocks"] == 2
    assert metrics["ocr_blocks"] == 0
    assert metrics["ocr_coverage"] == 1.0


def test_router_unsupported_file_type():
    """Verify router raises ValueError for unsupported extensions."""
    with pytest.raises(ValueError, match="Unsupported file type"):
        route_and_extract("file.xyz", "doc_x", "file.xyz")
