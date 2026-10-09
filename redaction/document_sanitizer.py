import os
import shutil
import re
import time
from typing import List, Dict, Any, Optional, Tuple
import docx
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
import fitz  # PyMuPDF

from models.schemas import DocumentBlock, PIIFinding, ValidationResult
from extractors.router import route_and_extract
from detection.pii_engine import detect_pii_in_document_blocks
from detection.entity_resolver import EntityResolver
from redaction.validator import is_block_uncertain
from utils.file_security import compute_file_sha256

REDACTION_TOKEN_PATTERN = re.compile(r'^\[[A-Z0-9_]+\]$')


def _get_sorted_replacements(findings: List[PIIFinding]) -> List[Tuple[str, str]]:
    """
    Deduplicate and sort findings by length descending (longest first)
    to prevent partial substring corruption.
    """
    rep_map = {}
    for f in sorted(findings, key=lambda x: len(x.text), reverse=True):
        if f.text not in rep_map:
            token = f.replacement_token if getattr(f, "replacement_token", None) else f"[{f.pii_type}_REDACTED]"
            rep_map[f.text] = token
    return list(rep_map.items())


def _finding_on_page(f: PIIFinding, page_idx: int) -> bool:
    """Check if finding belongs to the given 1-based page index."""
    loc = f.location.lower().strip()
    if f"page_{page_idx}_" in loc or loc == f"page_{page_idx}":
        return True
    if f"page {page_idx}" in loc:
        return True
    parts = loc.split("_")
    if len(parts) >= 2 and parts[0] == "page" and parts[1].isdigit():
        return int(parts[1]) == page_idx
    return False


def sanitize_docx(input_path: str, output_path: str, findings: List[PIIFinding]) -> Tuple[str, int]:
    """
    Structure-preserving DOCX sanitization:
    Preserves document styles, paragraphs, tables, headers, and footers
    while replacing sensitive PII mentions with resolved semantic tokens.
    """
    doc = docx.Document(input_path)
    replacements = _get_sorted_replacements(findings)
    applied_count = 0

    def _replace_in_paragraph(paragraph):
        nonlocal applied_count
        for orig_text, token in replacements:
            if orig_text in paragraph.text:
                replaced_in_run = False
                for run in paragraph.runs:
                    if orig_text in run.text:
                        run.text = run.text.replace(orig_text, token)
                        replaced_in_run = True
                        applied_count += 1
                if not replaced_in_run:
                    paragraph.text = paragraph.text.replace(orig_text, token)
                    applied_count += 1

    # 1. Main body paragraphs
    for p in doc.paragraphs:
        _replace_in_paragraph(p)

    # 2. Structured tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p)

    # 3. Section headers and footers
    for section in doc.sections:
        if section.header:
            for p in section.header.paragraphs:
                _replace_in_paragraph(p)
        if section.footer:
            for p in section.footer.paragraphs:
                _replace_in_paragraph(p)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path, applied_count


def sanitize_pptx(input_path: str, output_path: str, findings: List[PIIFinding]) -> Tuple[str, int]:
    """
    Structure-preserving PPTX sanitization:
    Preserves slide layouts, text shapes, recursive groups, and tables
    while replacing sensitive PII mentions with resolved semantic tokens.
    """
    prs = pptx.Presentation(input_path)
    replacements = _get_sorted_replacements(findings)
    applied_count = 0

    def _replace_in_text_frame(tf):
        nonlocal applied_count
        for p in tf.paragraphs:
            for orig_text, token in replacements:
                if orig_text in p.text:
                    replaced_in_run = False
                    for run in p.runs:
                        if orig_text in run.text:
                            run.text = run.text.replace(orig_text, token)
                            replaced_in_run = True
                            applied_count += 1
                    if not replaced_in_run:
                        p.text = p.text.replace(orig_text, token)
                        applied_count += 1

    def _process_shape(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                _process_shape(child)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if hasattr(cell, "text_frame"):
                        _replace_in_text_frame(cell.text_frame)
        elif hasattr(shape, "text_frame") and shape.text_frame:
            _replace_in_text_frame(shape.text_frame)

    for slide in prs.slides:
        for shape in slide.shapes:
            _process_shape(shape)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            _replace_in_text_frame(slide.notes_slide.notes_text_frame)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)
    return output_path, applied_count


def sanitize_pdf(input_path: str, output_path: str, findings: List[PIIFinding]) -> Tuple[str, int]:
    """
    Structure-preserving PDF sanitization using PyMuPDF:
    Applies visual redaction annotations over native text locations,
    scrubs underlying glyph/text streams, removes image pixels if applicable,
    and stamps replacement tokens.
    """
    doc = fitz.open(input_path)
    replacements = _get_sorted_replacements(findings)
    applied_count = 0

    # 1. Scrub document metadata if metadata findings exist
    has_meta_findings = any(f.location.lower().startswith("metadata") for f in findings)
    if has_meta_findings:
        try:
            doc.set_metadata({
                "title": "", "author": "", "subject": "", "keywords": "",
                "creator": "", "producer": "", "creationDate": "", "modDate": ""
            })
            applied_count += 1
        except Exception:
            pass

    # 2. Iterate across pages and apply redactions
    for page_num in range(len(doc)):
        page = doc[page_num]
        pw = max(1.0, float(page.rect.width))
        ph = max(1.0, float(page.rect.height))
        page_idx = page_num + 1

        redacted_rects: List[fitz.Rect] = []

        # A. Redact native text occurrences via search_for
        for orig_text, token in replacements:
            clean_search = orig_text.strip()
            if not clean_search or clean_search.startswith("<"):
                continue

            rects = page.search_for(clean_search)
            if not rects and len(clean_search) > 3:
                rects = page.search_for(clean_search, flags=fitz.TEXT_CASEFOLD)

            if rects:
                for r in rects:
                    page.add_redact_annot(r, text=token, fill=(0.95, 0.95, 0.95), text_color=(0, 0, 0))
                    redacted_rects.append(r)
                    applied_count += 1

        # B. Check for bbox-targeted findings on this specific page (OCR or Visual PII)
        page_findings = [f for f in findings if _finding_on_page(f, page_idx) and getattr(f, "bbox", None)]
        for f in page_findings:
            if f.bbox:
                bx0 = f.bbox["x0"] * pw
                by0 = f.bbox["y0"] * ph
                bx1 = f.bbox["x1"] * pw
                by1 = f.bbox["y1"] * ph
                r = fitz.Rect(bx0, by0, bx1, by1)

                already_covered = any(r.intersects(cov) for cov in redacted_rects)
                if not already_covered and r.width > 2 and r.height > 2:
                    token = f.replacement_token if getattr(f, "replacement_token", None) else f"[{f.pii_type}_REDACTED]"
                    page.add_redact_annot(r, text=token, fill=(0.95, 0.95, 0.95), text_color=(0, 0, 0))
                    redacted_rects.append(r)
                    applied_count += 1

        # Apply redactions including image pixels for scanned/image content
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    tmp_output = output_path + f".tmp_{os.getpid()}_{int(time.time()*1000)}"
    doc.save(tmp_output, garbage=4, deflate=True)
    doc.close()

    if os.path.exists(output_path):
        try:
            os.remove(output_path)
        except Exception:
            pass
    try:
        os.replace(tmp_output, output_path)
    except Exception:
        if os.path.exists(tmp_output):
            output_path = tmp_output

    return output_path, applied_count


def sanitize_txt(input_path: str, output_path: str, findings: List[PIIFinding]) -> Tuple[str, int]:
    """
    Sanitize plain text document using resolved semantic tokens.
    """
    text = ""
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            with open(input_path, "r", encoding=enc) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue

    replacements = _get_sorted_replacements(findings)
    applied_count = 0
    for orig_text, token in replacements:
        if orig_text in text:
            applied_count += text.count(orig_text)
            text = text.replace(orig_text, token)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)
    return output_path, applied_count


class DocumentSanitizer:
    """
    Central structure-preserving document sanitizer and coverage gate engine.
    Implements DETECT -> REDACT -> RE-SCAN VERIFY workflow.
    """
    def __init__(self, output_dir: str = "sanitized_output"):
        self.output_dir = output_dir
        self.last_applied_count: int = 0
        os.makedirs(self.output_dir, exist_ok=True)

    def sanitize(
        self,
        input_path: str,
        findings: List[PIIFinding],
        output_filename: Optional[str] = None,
        return_count: bool = False
    ) -> Any:
        """
        Dispatches format-specific structure-preserving sanitization.
        Returns output_path (str) by default, or (output_path, applied_count) if return_count=True.
        """
        base_name = os.path.basename(input_path)
        out_name = output_filename or f"sanitized_{base_name}"
        output_path = os.path.join(self.output_dir, out_name)

        ext = base_name.split(".")[-1].lower()

        if ext == "docx":
            out_p, count = sanitize_docx(input_path, output_path, findings)
        elif ext == "pptx":
            out_p, count = sanitize_pptx(input_path, output_path, findings)
        elif ext == "pdf":
            out_p, count = sanitize_pdf(input_path, output_path, findings)
        elif ext in ["txt", "log", "text"]:
            out_p, count = sanitize_txt(input_path, output_path, findings)
        else:
            shutil.copyfile(input_path, output_path)
            out_p, count = output_path, 0

        self.last_applied_count = count
        if return_count:
            return out_p, count
        return out_p

    def verify_and_rescan(
        self,
        sanitized_file_path: str,
        original_blocks: List[DocumentBlock],
        original_findings: List[PIIFinding],
        document_id: str,
        resolver: Optional[EntityResolver] = None,
        policy_engine: Optional[Any] = None,
        original_file_path: Optional[str] = None,
        applied_redactions_count: Optional[int] = None
    ) -> Tuple[ValidationResult, List[DocumentBlock]]:
        """
        Mandatory Phase 3 Re-Scan Verification Gate:
        1. Re-opens and re-extracts the generated sanitized file.
        2. Re-runs PII detection across all re-extracted blocks.
        3. Filters out valid semantic replacement tokens ([PERSON_001], etc.).
        4. Verifies cryptographic SHA difference from original when PII is present.
        5. Verifies finding-to-redaction linkage (applied_count >= len(original_findings)).
        6. Central fail-closed decision:
           - cryptographic SHA identical with PII -> BLOCKED
           - extraction uncertainty > 0 -> HUMAN_REVIEW
           - residual high-risk PII > 0 -> BLOCKED
           - residual low/medium PII > 0 -> HUMAN_REVIEW
           - incomplete redactions applied -> HUMAN_REVIEW
           - clean extraction & zero residual PII -> SANITIZED
        """
        effective_applied = applied_redactions_count if applied_redactions_count is not None else self.last_applied_count

        if not os.path.exists(sanitized_file_path):
            val = ValidationResult(
                document_id=document_id,
                is_safe=False,
                remaining_findings=[],
                message="Sanitized output file does not exist on disk.",
                requires_human_review=False,
                uncertain_blocks_count=0,
                pii_before=len(original_findings),
                pii_after=0,
                high_risk_residual=0,
                sanitized_file_path=sanitized_file_path,
                re_scanned_blocks_count=0,
                security_decision="BLOCKED",
                applied_redactions_count=0
            )
            return val, []

        try:
            # 1. Re-open and re-extract generated file
            re_extracted_blocks = route_and_extract(
                file_path=sanitized_file_path,
                document_id=f"{document_id}_rescan",
                original_file_name=os.path.basename(sanitized_file_path)
            )
        except Exception as e:
            val = ValidationResult(
                document_id=document_id,
                is_safe=False,
                remaining_findings=[],
                message=f"Sanitized file failed re-extraction validation: {str(e)}",
                requires_human_review=False,
                uncertain_blocks_count=0,
                pii_before=len(original_findings),
                pii_after=0,
                high_risk_residual=0,
                sanitized_file_path=sanitized_file_path,
                re_scanned_blocks_count=0,
                security_decision="BLOCKED",
                applied_redactions_count=effective_applied
            )
            return val, []

        # 2. Check for extraction uncertainty
        orig_uncertain = sum(1 for b in original_blocks if is_block_uncertain(b))
        rescan_uncertain = sum(1 for b in re_extracted_blocks if is_block_uncertain(b))
        total_uncertain = max(orig_uncertain, rescan_uncertain)

        # 3. Re-run PII detection on re-extracted blocks
        raw_rescan_findings = detect_pii_in_document_blocks(re_extracted_blocks, resolver=resolver)

        # 4. Filter out legitimate semantic replacement tokens
        residual_findings = []
        for f in raw_rescan_findings:
            if REDACTION_TOKEN_PATTERN.match(f.text.strip()):
                continue
            residual_findings.append(f)

        # 5. Check for high-risk residual findings
        high_risk_residual = [
            f for f in residual_findings
            if f.severity in ["CRITICAL", "HIGH"] or f.confidence >= 0.85
        ]

        # 6. Cryptographic SHA check against original
        sha_identical_error = False
        if original_file_path and os.path.exists(original_file_path):
            try:
                orig_sha = compute_file_sha256(original_file_path)
                sanit_sha = compute_file_sha256(sanitized_file_path)
                if len(original_findings) > 0 and orig_sha == sanit_sha:
                    sha_identical_error = True
            except Exception:
                pass

        # 7. Redaction application linkage check
        incomplete_redactions = False
        if len(original_findings) > 0 and effective_applied < len(original_findings):
            incomplete_redactions = True

        # 8. Central Fail-Closed Decision Gate
        if sha_identical_error:
            is_safe = False
            requires_human_review = False
            decision = "BLOCKED"
            message = (
                "SECURITY ERROR: Sanitized artifact is cryptographically identical to original artifact despite detected PII."
            )
        elif total_uncertain > 0:
            is_safe = False
            requires_human_review = True
            decision = "HUMAN_REVIEW"
            message = (
                f"Extraction uncertainty detected ({total_uncertain} uninspected/failed blocks). "
                f"Fail-closed policy requires HUMAN REVIEW."
            )
        elif len(high_risk_residual) > 0:
            is_safe = False
            requires_human_review = False
            decision = "BLOCKED"
            message = (
                f"Residual PII detected ({len(high_risk_residual)} high-risk findings) in post-redaction re-scan. "
                f"Release denied."
            )
        elif len(residual_findings) > 0:
            is_safe = False
            requires_human_review = True
            decision = "HUMAN_REVIEW"
            message = (
                f"Potential low/medium residual PII detected ({len(residual_findings)} findings) in re-scan. "
                f"Manual review required."
            )
        elif incomplete_redactions:
            is_safe = False
            requires_human_review = True
            decision = "HUMAN_REVIEW"
            message = (
                f"Incomplete redaction: expected {len(original_findings)} findings to be redacted "
                f"but only {effective_applied} were applied. Manual review required."
            )
        else:
            is_safe = True
            requires_human_review = False
            decision = "SANITIZED"
            message = (
                f"Zero residual PII detected and 100% extraction coverage verified. "
                f"Document safe for downstream AI release."
            )

        val = ValidationResult(
            document_id=document_id,
            is_safe=is_safe,
            remaining_findings=residual_findings,
            message=message,
            requires_human_review=requires_human_review,
            uncertain_blocks_count=total_uncertain,
            pii_before=len(original_findings),
            pii_after=len(residual_findings),
            high_risk_residual=len(high_risk_residual),
            sanitized_file_path=sanitized_file_path,
            re_scanned_blocks_count=len(re_extracted_blocks),
            security_decision=decision,
            applied_redactions_count=effective_applied
        )

        # 9. Apply optional PolicyEngine Gate Evaluation
        if policy_engine is not None:
            val = policy_engine.evaluate_rescan_gate(val)

        return val, re_extracted_blocks


default_document_sanitizer = DocumentSanitizer()
